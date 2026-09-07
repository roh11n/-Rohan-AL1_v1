"""Parse KB uploads (CSV/Excel/JSON/PDF/DOCX/MD) into text chunks."""
import io
import json
import logging
import re
from typing import Any

logger = logging.getLogger(__name__)


def _split_text(text: str, chunk_size: int = 800, overlap: int = 100) -> list[str]:
    text = re.sub(r"\s+", " ", text or "").strip()
    if not text:
        return []
    chunks = []
    i = 0
    while i < len(text):
        end = min(len(text), i + chunk_size)
        chunks.append(text[i:end])
        if end == len(text):
            break
        i = end - overlap
    return chunks


def parse_upload(filename: str, content: bytes) -> list[str]:
    """Return a list of text chunks extracted from the uploaded file."""
    name = (filename or "").lower()
    try:
        if name.endswith(".csv"):
            text = content.decode("utf-8", errors="ignore")
            return _split_text(text)
        if name.endswith(".json"):
            try:
                data = json.loads(content.decode("utf-8", errors="ignore"))
            except Exception:
                data = content.decode("utf-8", errors="ignore")
            text = json.dumps(data, indent=2) if not isinstance(data, str) else data
            return _split_text(text)
        if name.endswith(".xlsx") or name.endswith(".xls"):
            try:
                import openpyxl  # type: ignore
                wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
                lines = []
                for ws in wb.worksheets:
                    for row in ws.iter_rows(values_only=True):
                        line = " | ".join("" if c is None else str(c) for c in row)
                        if line.strip():
                            lines.append(line)
                text = "\n".join(lines)
                return _split_text(text)
            except Exception as e:
                logger.warning("xlsx parse failed: %s", e)
                return []
        if name.endswith(".pdf"):
            try:
                from PyPDF2 import PdfReader  # type: ignore
                reader = PdfReader(io.BytesIO(content))
                text = "\n".join((p.extract_text() or "") for p in reader.pages)
                return _split_text(text)
            except Exception as e:
                logger.warning("pdf parse failed: %s", e)
                return []
        if name.endswith(".docx"):
            try:
                import docx  # type: ignore
                d = docx.Document(io.BytesIO(content))
                text = "\n".join(p.text for p in d.paragraphs)
                return _split_text(text)
            except Exception as e:
                logger.warning("docx parse failed: %s", e)
                return []
        # md / txt / others
        text = content.decode("utf-8", errors="ignore")
        return _split_text(text)
    except Exception as e:
        logger.warning("KB parse failed for %s: %s", filename, e)
        return []


def content_summary(chunks: list[str]) -> str:
    if not chunks:
        return "(empty)"
    joined = " ".join(chunks[:2])
    return joined[:280] + ("..." if len(joined) > 280 else "")
